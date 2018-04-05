import docx,openpyxl


#创建带有默认样式的word文件
def creatDoc(templeDest):
    doc = docx.Document(templeDest)
    return doc

#打开excle文件
def openQsBank(filename):
    wb = openpyxl.load_workbook(filename)
    acSheet = wb.get_active_sheet()
    return acSheet

#获取单元格的数值
def getCellValue(sheet, stRow, stColumn):
    return sheet.cell(row= stRow, column=stColumn).value

#写入单元格的数值
def setCellValue(sheet,stRow,stColumn,s):
    sheet.cell(row=stRow,column=stColumn).value = s
    return True

#确定题量，返回各类试题所在行数
def calQueNum(useSheet):
    singleChoice = []
    multiChoice = []
    judgeChoice = []
    jQuiz = []
    for i in range(1,500):
        if getCellValue(useSheet,i,1) == r'单选题' :
            singleChoice.append(i)
        elif getCellValue(useSheet,i,1) == r'多选题' :
            multiChoice.append(i)
        elif getCellValue(useSheet,i,1) == r'判断题' :
            judgeChoice.append(i)
        elif getCellValue(useSheet,i,1) == r'简答题':
            jQuiz.append(i)
    return singleChoice,multiChoice,judgeChoice,jQuiz

